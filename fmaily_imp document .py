

import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from pathlib import Path

FILE = Path(__file__).with_name("family_organizer.xlsx")
BG = "#F3F7FC"
WHITE = "#FFFFFF"
NAVY = "#253F9B"
BLUE = "#2878E8"
GREEN = "#12A889"
RED = "#F04444"
PURPLE = "#5B61E8"
TEXT = "#172B4D"
MUTED = "#64748B"


def setup_excel():
    if not FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Contacts"
        ws.append(["ID", "Name", "Phone", "Relationship", "Type", "Primary"])
        ws = wb.create_sheet("Documents")
        ws.append(["ID", "Document Name", "Category", "Number", "Location", "Notes"])
        wb.save(FILE)
    else:
        wb = load_workbook(FILE)
        if "Contacts" not in wb.sheetnames:
            ws = wb.create_sheet("Contacts")
            ws.append(["ID", "Name", "Phone", "Relationship", "Type", "Primary"])
        if "Documents" not in wb.sheetnames:
            ws = wb.create_sheet("Documents")
            ws.append(["ID", "Document Name", "Category", "Number", "Location", "Notes"])
        wb.save(FILE)


def rows(sheet):
    setup_excel()
    wb = load_workbook(FILE)
    return [list(r) for r in wb[sheet].iter_rows(min_row=2, values_only=True)
            if any(v is not None for v in r)]


def save_record(sheet, values, record_id=None):
    wb = load_workbook(FILE)
    ws = wb[sheet]
    if record_id is None:
        ids = [r[0] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True)
               if isinstance(r[0], int)]
        ws.append([max(ids + [0]) + 1] + values)
    else:
        for n in range(2, ws.max_row + 1):
            if ws.cell(n, 1).value == record_id:
                for col, val in enumerate(values, start=2):
                    ws.cell(n, col).value = val
                break
    wb.save(FILE)


def delete_record(sheet, record_id):
    wb = load_workbook(FILE)
    ws = wb[sheet]
    for n in range(2, ws.max_row + 1):
        if ws.cell(n, 1).value == record_id:
            ws.delete_rows(n)
            break
    wb.save(FILE)


class Organizer:
    def __init__(self, root):
        self.root = root
        root.title("Family Emergency Contact & Document Organizer")
        root.geometry("1200x700")
        root.minsize(1000, 620)
        root.configure(bg=BG)
        self.contact_id = None
        self.document_id = None

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Treeview", background=WHITE, fieldbackground=WHITE,
                        foreground=TEXT, rowheight=30, font=("Segoe UI", 9))
        style.configure("Treeview.Heading", background="#EDF3FC",
                        foreground=TEXT, font=("Segoe UI", 9, "bold"), relief="flat")
        style.map("Treeview", background=[("selected", "#B9D9FF")])

        self.sidebar = tk.Frame(root, bg=WHITE, width=220,
                                 highlightbackground="#D7E5F7", highlightthickness=1)
        self.sidebar.pack(side="left", fill="y")
        self.sidebar.pack_propagate(False)
        tk.Label(self.sidebar, text="✚  FAMILY CARE", bg=WHITE, fg=NAVY,
                 font=("Segoe UI", 15, "bold"), pady=25).pack(fill="x")
        self.nav = {}
        self.nav_button("⌂", "Home", self.home)
        self.nav_button("☎", "Emergency Contacts", self.contacts)
        self.nav_button("▣", "Important Documents", self.documents)
        tk.Label(self.sidebar, text="Keep your family safe.\\nData saved in Excel.",
                 bg=WHITE, fg=MUTED, font=("Segoe UI", 9), justify="left").pack(
                     side="bottom", anchor="w", padx=20, pady=24)

        self.main = tk.Frame(root, bg=BG)
        self.main.pack(side="left", fill="both", expand=True, padx=18, pady=16)
        self.home()

    def nav_button(self, icon, label, command):
        b = tk.Button(self.sidebar, text=f"{icon}   {label}", anchor="w",
                      command=command, padx=18, pady=14, bd=0, bg=WHITE,
                      fg=TEXT, activebackground="#E8F2FF", font=("Segoe UI", 10))
        b.pack(fill="x", padx=10, pady=4)
        self.nav[label] = b

    def active_nav(self, label):
        for name, b in self.nav.items():
            b.configure(bg="#E8F2FF" if name == label else WHITE,
                        fg=BLUE if name == label else TEXT,
                        font=("Segoe UI", 10, "bold" if name == label else "normal"))

    def clear(self):
        for w in self.main.winfo_children():
            w.destroy()

    def heading(self, icon, title, subtitle):
        bar = tk.Frame(self.main, bg=BG)
        bar.pack(fill="x", pady=(0, 16))
        tk.Label(bar, text=icon, bg=NAVY, fg=WHITE, width=3, pady=7,
                 font=("Segoe UI", 20, "bold")).pack(side="left", padx=(0, 12))
        t = tk.Frame(bar, bg=BG)
        t.pack(side="left", fill="x", expand=True)
        tk.Label(t, text=title, bg=BG, fg=NAVY,
                 font=("Segoe UI", 19, "bold")).pack(anchor="w")
        tk.Label(t, text=subtitle, bg=BG, fg=MUTED,
                 font=("Segoe UI", 9)).pack(anchor="w", pady=(3, 0))

    def button(self, parent, text, command, color):
        return tk.Button(parent, text=text, command=command, bg=color, fg=WHITE,
                         activebackground=color, activeforeground=WHITE, bd=0,
                         padx=12, pady=8, cursor="hand2",
                         font=("Segoe UI", 9, "bold"))

    def home(self):
        self.clear()
        self.active_nav("Home")
        hero = tk.Frame(self.main, bg=NAVY, height=180)
        hero.pack(fill="x", pady=(0, 18))
        hero.pack_propagate(False)
        tk.Label(hero, text="👨‍👩‍👧", bg=NAVY, fg=WHITE,
                 font=("Segoe UI Emoji", 30)).pack(anchor="w", padx=24, pady=(14, 0))
        tk.Label(hero, text="Family Emergency Contact and\\nImportant Document Organizer",
                 bg=NAVY, fg=WHITE, justify="left",
                 font=("Segoe UI", 19, "bold")).pack(anchor="w", padx=24)
        tk.Label(hero, text="Keep your family safe and your important details organized.",
                 bg=NAVY, fg="#E1E9FF", font=("Segoe UI", 10)).pack(anchor="w", padx=24, pady=8)

        counts = tk.Frame(self.main, bg=BG)
        counts.pack(fill="x", pady=(0, 18))
        self.stat(counts, "☎", "Emergency Contacts", len(rows("Contacts")), "#FFF0F1", RED, 0)
        self.stat(counts, "▣", "Important Documents", len(rows("Documents")), "#EAF3FF", BLUE, 1)

        actions = tk.Frame(self.main, bg=BG)
        actions.pack(fill="x", pady=(0, 18))
        self.button(actions, "☎   Manage Emergency Contacts", self.contacts, GREEN).pack(
            side="left", fill="x", expand=True, padx=(0, 8), ipady=7)
        self.button(actions, "▣   Manage Important Documents", self.documents, PURPLE).pack(
            side="left", fill="x", expand=True, padx=(8, 0), ipady=7)
        info = tk.Frame(self.main, bg="#EAF3FF", padx=16, pady=16)
        info.pack(fill="x")
        tk.Label(info, text="ⓘ", bg="#EAF3FF", fg=BLUE,
                 font=("Segoe UI", 18, "bold")).pack(side="left", padx=(0, 12))
        tk.Label(info, text="Your details are saved in family_organizer.xlsx in the same folder.\\n"
                            "Keep this file safe and private.",
                 bg="#EAF3FF", fg="#36557C", justify="left",
                 font=("Segoe UI", 10)).pack(side="left")

    def stat(self, parent, icon, title, count, color, accent, col):
        parent.grid_columnconfigure(col, weight=1)
        f = tk.Frame(parent, bg=color, padx=20, pady=18,
                     highlightbackground="#DCE8F8", highlightthickness=1)
        f.grid(row=0, column=col, sticky="nsew", padx=(0, 8) if col == 0 else (8, 0))
        tk.Label(f, text=icon, bg=color, fg=accent,
                 font=("Segoe UI", 24, "bold")).pack(anchor="w")
        tk.Label(f, text=title, bg=color, fg=TEXT,
                 font=("Segoe UI", 11, "bold")).pack(anchor="w", pady=(7, 2))
        tk.Label(f, text=str(count), bg=color, fg=TEXT,
                 font=("Segoe UI", 25, "bold")).pack(anchor="w")
        tk.Label(f, text="Total saved", bg=color, fg=MUTED,
                 font=("Segoe UI", 9)).pack(anchor="w")

    def entry(self, parent, label, width=None):
        tk.Label(parent, text=label, bg=WHITE, fg=TEXT,
                 font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(7, 3))
        e = tk.Entry(parent, font=("Segoe UI", 9), relief="solid", bd=1)
        e.pack(fill="x", ipady=6)
        return e

    def layout_form_table(self, title):
        body = tk.Frame(self.main, bg=BG)
        body.pack(fill="both", expand=True)
        body.grid_columnconfigure(0, weight=0, minsize=300)
        body.grid_columnconfigure(1, weight=1)
        body.grid_rowconfigure(0, weight=1)
        form = tk.Frame(body, bg=WHITE, padx=16, pady=14,
                        highlightbackground="#DFE9F5", highlightthickness=1)
        form.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        table = tk.Frame(body, bg=WHITE, padx=10, pady=12,
                         highlightbackground="#DFE9F5", highlightthickness=1)
        table.grid(row=0, column=1, sticky="nsew")
        tk.Label(form, text=title, bg=WHITE, fg=NAVY,
                 font=("Segoe UI", 12, "bold")).pack(anchor="w", pady=(0, 8))
        return form, table

    def contacts(self):
        self.clear()
        self.active_nav("Emergency Contacts")
        self.heading("☎", "Emergency Contacts",
                     "Add, view, update or delete your family emergency contacts.")
        form, table = self.layout_form_table("Contact Details")
        self.c_name = self.entry(form, "Name")
        self.c_phone = self.entry(form, "Phone Number")
        self.c_relation = self.entry(form, "Relationship")
        tk.Label(form, text="Emergency Type", bg=WHITE, fg=TEXT,
                 font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(8, 4))
        self.c_type = ttk.Combobox(form, values=["Medical", "Family", "Police", "Fire", "Other"],
                                   state="readonly")
        self.c_type.set("Medical")
        self.c_type.pack(fill="x")
        self.c_primary = tk.BooleanVar()
        tk.Checkbutton(form, text="Primary Contact", variable=self.c_primary,
                       bg=WHITE, font=("Segoe UI", 9)).pack(anchor="w", pady=8)
        self.row_buttons(form, self.add_contact, self.update_contact,
                         self.delete_contact, self.clear_contact)
        cols = ("ID", "Name", "Phone", "Relationship", "Type", "Primary")
        self.ctree = ttk.Treeview(table, columns=cols, show="headings")
        for col, width in zip(cols, [38, 120, 100, 100, 85, 65]):
            self.ctree.heading(col, text=col)
            self.ctree.column(col, width=width, anchor="center")
        self.ctree.pack(fill="both", expand=True)
        self.ctree.bind("<<TreeviewSelect>>", self.pick_contact)
        self.refresh("Contacts", self.ctree)

    def row_buttons(self, form, add, update, delete, clear):
        f = tk.Frame(form, bg=WHITE)
        f.pack(fill="x", pady=(12, 0))
        for label, cmd, color in [("Add", add, GREEN), ("Update", update, BLUE),
                                  ("Delete", delete, RED), ("Clear", clear, "#8293A8")]:
            self.button(f, label, cmd, color).pack(side="left", padx=(0, 5))

    def refresh(self, sheet, tree):
        for item in tree.get_children():
            tree.delete(item)
        for r in rows(sheet):
            r = r + [""] * (6 - len(r))
            tree.insert("", "end", values=r)

    def pick_contact(self, _event=None):
        selected = self.ctree.selection()
        if not selected:
            return
        v = self.ctree.item(selected[0], "values")
        self.contact_id = int(v[0])
        for e, value in zip([self.c_name, self.c_phone, self.c_relation], v[1:4]):
            e.delete(0, "end")
            e.insert(0, value or "")
        self.c_type.set(v[4] or "Medical")
        self.c_primary.set(v[5] == "Yes")

    def clear_contact(self):
        self.contact_id = None
        for e in [self.c_name, self.c_phone, self.c_relation]:
            e.delete(0, "end")
        self.c_type.set("Medical")
        self.c_primary.set(False)

    def contact_values(self):
        name, phone, rel = self.c_name.get().strip(), self.c_phone.get().strip(), self.c_relation.get().strip()
        if not name or not phone or not rel:
            messagebox.showwarning("Missing Details", "Please fill in name, phone number and relationship.")
            return None
        if not phone.isdigit() or not 7 <= len(phone) <= 15:
            messagebox.showwarning("Invalid Phone", "Enter a valid phone number (7 to 15 digits).")
            return None
        return [name, phone, rel, self.c_type.get(), "Yes" if self.c_primary.get() else "No"]

    def add_contact(self):
        v = self.contact_values()
        if v:
            save_record("Contacts", v)
            self.refresh("Contacts", self.ctree)
            self.clear_contact()
            messagebox.showinfo("Saved", "Emergency contact added.")

    def update_contact(self):
        if self.contact_id is None:
            messagebox.showwarning("Select Contact", "Select a contact from the table first.")
            return
        v = self.contact_values()
        if v:
            save_record("Contacts", v, self.contact_id)
            self.refresh("Contacts", self.ctree)
            self.clear_contact()
            messagebox.showinfo("Updated", "Contact updated.")

    def delete_contact(self):
        if self.contact_id is None:
            messagebox.showwarning("Select Contact", "Select a contact from the table first.")
            return
        if messagebox.askyesno("Confirm Delete", "Are you sure you want to delete this contact?"):
            delete_record("Contacts", self.contact_id)
            self.refresh("Contacts", self.ctree)
            self.clear_contact()

    def documents(self):
        self.clear()
        self.active_nav("Important Documents")
        self.heading("▣", "Important Documents",
                     "Keep details of your important documents safe and organized.")
        form, table = self.layout_form_table("Document Details")
        self.d_name = self.entry(form, "Document Name")
        tk.Label(form, text="Category", bg=WHITE, fg=TEXT,
                 font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(8, 4))
        self.d_category = ttk.Combobox(form, values=["Identity", "Insurance", "Banking",
                                      "Medical", "Education", "Property", "Other"],
                                       state="readonly")
        self.d_category.set("Identity")
        self.d_category.pack(fill="x")
        self.d_number = self.entry(form, "Document Number")
        self.d_location = self.entry(form, "Location")
        tk.Label(form, text="Notes", bg=WHITE, fg=TEXT,
                 font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(8, 4))
        self.d_notes = tk.Text(form, height=3, font=("Segoe UI", 9), relief="solid", bd=1)
        self.d_notes.pack(fill="x")
        self.row_buttons(form, self.add_document, self.update_document,
                         self.delete_document, self.clear_document)
        cols = ("ID", "Document Name", "Category", "Number", "Location", "Notes")
        self.dtree = ttk.Treeview(table, columns=cols, show="headings")
        for col, width in zip(cols, [38, 125, 85, 110, 105, 110]):
            self.dtree.heading(col, text=col)
            self.dtree.column(col, width=width, anchor="w" if col in ("Document Name", "Location", "Notes") else "center")
        self.dtree.pack(fill="both", expand=True)
        self.dtree.bind("<<TreeviewSelect>>", self.pick_document)
        self.refresh("Documents", self.dtree)

    def pick_document(self, _event=None):
        selected = self.dtree.selection()
        if not selected:
            return
        v = self.dtree.item(selected[0], "values")
        self.document_id = int(v[0])
        for e, value in zip([self.d_name, self.d_number, self.d_location], [v[1], v[3], v[4]]):
            e.delete(0, "end")
            e.insert(0, value or "")
        self.d_category.set(v[2] or "Identity")
        self.d_notes.delete("1.0", "end")
        self.d_notes.insert("1.0", v[5] or "")

    def clear_document(self):
        self.document_id = None
        for e in [self.d_name, self.d_number, self.d_location]:
            e.delete(0, "end")
        self.d_category.set("Identity")
        self.d_notes.delete("1.0", "end")

    def document_values(self):
        name, number, location = self.d_name.get().strip(), self.d_number.get().strip(), self.d_location.get().strip()
        notes = self.d_notes.get("1.0", "end").strip()
        if not name or not number or not location:
            messagebox.showwarning("Missing Details", "Please fill in document name, number and location.")
            return None
        return [name, self.d_category.get(), number, location, notes]

    def add_document(self):
        v = self.document_values()
        if v:
            save_record("Documents", v)
            self.refresh("Documents", self.dtree)
            self.clear_document()
            messagebox.showinfo("Saved", "Document details added.")

    def update_document(self):
        if self.document_id is None:
            messagebox.showwarning("Select Document", "Select a document from the table first.")
            return
        v = self.document_values()
        if v:
            save_record("Documents", v, self.document_id)
            self.refresh("Documents", self.dtree)
            self.clear_document()
            messagebox.showinfo("Updated", "Document details updated.")

    def delete_document(self):
        if self.document_id is None:
            messagebox.showwarning("Select Document", "Select a document from the table first.")
            return
        if messagebox.askyesno("Confirm Delete", "Are you sure you want to delete this document?"):
            delete_record("Documents", self.document_id)
            self.refresh("Documents", self.dtree)
            self.clear_document()


if __name__ == "__main__":
    setup_excel()
    root = tk.Tk()
    Organizer(root)
    root.mainloop()
    
path = Path("/mnt/data/family_emergency_document_organizer.py")
path.write_text(code, encoding="utf-8")
print(f"Created Python file: {path}")
